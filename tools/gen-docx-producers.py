#!/usr/bin/env python3
"""Generate crates/docx-text/tests/fixtures/docx-producers.v1.json (noevia#981).

Producer compatibility: DOCX files written by real producers, from synthetic content only, must
be extracted by the docx-text binary with text byte-identical to the Python reference. The
binary's ZIP rules are deliberately stricter than zipfile; this set is the evidence that real
files are not caught by them. Producers (whichever are installed; the file records versions):

  python-docx            (uv run --with python-docx; images, tables, headers/footers, comments,
                          tracked changes injected as w:ins/w:del, core properties, a ~3 MB file)
  pandoc                 (markdown -> docx: tables, footnotes, image, unicode/RTL)
  LibreOffice headless   (soffice --convert-to docx from HTML, and re-saving the other
                          producers' files: comments and tracked changes survive the round trip)
  macOS textutil         (-convert docx from HTML and RTF)
  Java ZipOutputStream   (tools/JavaRepack.java: data descriptors, zero local sizes, as Java
                          exporters such as Google Docs write them)
  zip re-packers         (Info-ZIP zip, ditto -c -k, Python zipfile, each re-packing a docx
                          with an extra non-ASCII member name; Info-ZIP with -z adds a comment)

    cargo build --release -p docx-text-cli
    uv run --python 3.12 --with python-docx tools/gen-docx-producers.py \\
        --ocr <noevia-services>/ocr --bin target/release/docx-text

Each record holds the docx (base64), the producer and its version, and both results. Every
producer file must be extracted by Python and by the binary with identical text, except files
listed in EXPECTED_REFUSALS (a zip-tool output the binary refuses on purpose, with the reason).
"""
from __future__ import annotations

import argparse
import base64
import importlib
import io
import json
import os
import random
import shutil
import struct
import subprocess
import sys
import tempfile
import zipfile
import zlib
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
gen = importlib.import_module("gen-docx-text")

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "crates" / "docx-text" / "tests" / "fixtures" / "docx-producers.v1.json"
SOFFICE = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"

# Archive comments: Info-ZIP's `zip -z` is the only common tool that writes one, and it is an
# explicit user action. The binary refuses comments (they can hide a second archive), so this
# output is expected to be refused; see the crate docs.
EXPECTED_REFUSALS = {
    "zip-infozip-with-comment": "archive comment (zip -z): refused on purpose, a comment can carry a second archive",
}

TEXT_EN = ["Synthetic quarterly note for an invented company.",
           "Revenue rose in the fictional north region; costs fell slightly.",
           "This paragraph exists only to exercise the extractor."]
TEXT_INTL = ["Grüße aus Köln — naïve café, Œuvre, ß ẞ.", "中文测试段落，日本語の文章、한국어 문장.",
             "مرحبا، هذه فقرة تجريبية.", "שלום, זו פסקה לבדיקה.", "Emoji 🙂 and math ∑ √ ∞."]


def png(w: int, h: int, seed: int, noise: bool) -> bytes:
    rng = random.Random(seed)
    rows = []
    for y in range(h):
        if noise:
            row = bytes(rng.getrandbits(8) for _ in range(w * 3))
        else:
            row = bytes(((x * 7 + y * 3) % 256) for x in range(w)) * 3
        rows.append(b"\0" + row[: w * 3])
    raw = zlib.compress(b"".join(rows), 6)

    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d))
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + chunk(b"IDAT", raw) + chunk(b"IEND", b"")


def python_docx(tmp: Path) -> dict[str, bytes]:
    import docx
    from docx.shared import Inches
    out = {}
    img = tmp / "chart.png"
    img.write_bytes(png(120, 80, 1, False))

    def base(d):
        d.core_properties.author = "Synthetic Author"
        d.core_properties.title = "Synthetic"
        sec = d.sections[0]
        sec.header.paragraphs[0].text = "Synthetic header text"
        sec.footer.paragraphs[0].text = "Synthetic footer text"
        d.add_heading("Synthetic report", 0)
        for t in TEXT_EN + TEXT_INTL:
            d.add_paragraph(t)
        p = d.add_paragraph("Bold, ")
        p.add_run("italic").italic = True
        p.add_run(" and a tab\there, a break")
        p.add_run().add_break()
        p.add_run("after break.")
        table = d.add_table(rows=3, cols=3)
        for r in range(3):
            for c in range(3):
                table.cell(r, c).text = f"r{r}c{c} ✓"
        table.cell(0, 0).merge(table.cell(0, 1))
        d.add_picture(str(img), width=Inches(1.5))
        d.add_paragraph("Bulleted item", style="List Bullet")
        d.add_paragraph("Numbered item", style="List Number")
        return d

    d = base(docx.Document())
    buf = io.BytesIO(); d.save(buf); out["python-docx-basic"] = buf.getvalue()

    d = base(docx.Document())
    para = d.add_paragraph("Commented text.")
    if hasattr(d, "add_comment"):
        d.add_comment(para.runs, text="Synthetic reviewer comment", author="Reviewer", initials="RV")
    # Tracked changes as Word writes them (python-docx has no API for them).
    from docx.oxml import parse_xml
    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    d.element.body.insert(len(d.element.body) - 1, parse_xml(
        f'<w:p {W}><w:r><w:t xml:space="preserve">Kept </w:t></w:r>'
        f'<w:ins w:id="901" w:author="Reviewer" w:date="2020-01-01T00:00:00Z"><w:r><w:t>inserted</w:t></w:r></w:ins>'
        f'<w:del w:id="902" w:author="Reviewer" w:date="2020-01-01T00:00:00Z"><w:r><w:delText>deleted</w:delText></w:r></w:del></w:p>'))
    buf = io.BytesIO(); d.save(buf); out["python-docx-comments-tracked-changes"] = buf.getvalue()

    d = base(docx.Document())
    big = tmp / "noise.png"
    big.write_bytes(png(1000, 900, 2, True))
    d.add_picture(str(big), width=Inches(5))
    rng = random.Random(3)
    for i in range(1500):
        d.add_paragraph(f"Paragraph {i}: " + " ".join(rng.choice(gen.WORDS) for _ in range(12)))
    buf = io.BytesIO(); d.save(buf); out["python-docx-large"] = buf.getvalue()

    # Nested tables as Word users build them: 25 levels, a picture in each cell.
    d = docx.Document()
    cell = d.add_table(rows=1, cols=2).cell(0, 0)
    for lvl in range(25):
        cell.paragraphs[0].text = f"level {lvl}"
        cell.paragraphs[0].add_run().add_picture(str(img), width=Inches(0.3))
        t = cell.add_table(rows=1, cols=2)
        t.cell(0, 1).text = f"side {lvl}"
        cell = t.cell(0, 0)
    cell.paragraphs[0].text = "innermost"
    buf = io.BytesIO(); d.save(buf); out["python-docx-nested-tables"] = buf.getvalue()
    return out


MD = """---
title: Synthetic pandoc document
author: Synthetic Author
---

# Heading one

{intl}

Text with a footnote.[^1] And **bold**, *italic*, `code`.

| Item | Qty | Note |
|------|----:|------|
| Widget | 3 | ✓ grün |
| Gadget | 12 | 中文 |

![Synthetic chart](chart.png)

- bullet one
- bullet two

> A quoted synthetic line.

[^1]: The synthetic footnote.
""".format(intl="\n\n".join(TEXT_EN + TEXT_INTL))

HTML = """<html><head><meta charset="utf-8"><title>Synthetic</title></head><body>
<h1>Synthetic HTML document</h1>
{paras}
<table border="1"><tr><th>Item</th><th>Qty</th></tr><tr><td>Widget</td><td>3</td></tr><tr><td>Gadget ✓</td><td>12</td></tr></table>
<p>Image: <img src="chart.png" width="120" height="80"></p>
<ul><li>bullet one</li><li>bullet two</li></ul>
</body></html>""".format(paras="\n".join(f"<p>{t}</p>" for t in TEXT_EN + TEXT_INTL))

RTF = r"""{\rtf1\ansi\ansicpg1252\deff0{\fonttbl{\f0 Helvetica;}}
\f0\fs24 Synthetic RTF document.\par
Second paragraph with caf\'e9 and na\'efve.\par
\b Bold\b0  and \i italic\i0 .\par
\trowd\cellx2000\cellx4000 Item\cell Qty\cell\row
\trowd\cellx2000\cellx4000 Widget\cell 3\cell\row
}"""


def version(cmd: list[str]) -> str:
    try:
        out = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
        return (out.stdout or out.stderr).strip().splitlines()[0]
    except (OSError, IndexError, subprocess.TimeoutExpired):
        return "unavailable"


def soffice_convert(src: Path, tmp: Path, name: str) -> bytes | None:
    if not Path(SOFFICE).exists():
        return None
    outdir = tmp / f"lo-{name}"
    outdir.mkdir()
    profile = tmp / "lo-profile"
    subprocess.run([SOFFICE, f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
                    "--convert-to", "docx:MS Word 2007 XML", "--outdir", str(outdir), str(src)],
                   capture_output=True, timeout=300, check=True)
    return next(outdir.glob("*.docx")).read_bytes()


def repack(tmp: Path, name: str, docx_bytes: bytes, tool: str) -> bytes:
    """Unpack a docx, add a member with a non-ASCII name, re-pack it with a zip tool."""
    src = tmp / f"unz-{name}"
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        z.extractall(src)
    (src / "customXml").mkdir(exist_ok=True)
    (src / "customXml" / "notiz-ä-日本.txt").write_text("synthetic attachment")
    out = tmp / f"{name}.docx"
    if tool == "infozip":
        subprocess.run(["zip", "-q", "-X", "-r", str(out), "."], cwd=src, check=True)
    elif tool == "infozip-comment":
        subprocess.run(["zip", "-q", "-X", "-r", "-z", str(out), "."], cwd=src, check=True,
                       input=b"Synthetic archive comment\n.\n")
    elif tool == "ditto":
        subprocess.run(["ditto", "-c", "-k", "--norsrc", str(src), str(out)], check=True)
    else:
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            for f in sorted(src.rglob("*")):
                if f.is_file():
                    z.write(f, f.relative_to(src).as_posix())
    return out.read_bytes()


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--ocr", required=True)
    ap.add_argument("--bin", required=True)
    ap.add_argument("--out", default=str(OUT))
    a = ap.parse_args()
    sys.path.insert(0, a.ocr)
    mod = importlib.import_module("docx_text")
    fn = getattr(mod, "extract_docx_py", None) or mod.extract_docx
    import docx as _docx

    versions = {
        "python-docx": getattr(_docx, "__version__", "unknown"),
        "pandoc": version(["pandoc", "--version"]),
        "libreoffice": version([SOFFICE, "--version"]),
        "textutil": "macOS " + version(["sw_vers", "-productVersion"]),
        "infozip": next((l.strip() for l in subprocess.run(["zip", "-v"], capture_output=True, text=True).stdout.splitlines()
                         if l.startswith("This is Zip")), "unknown"),
        "ditto": "macOS " + version(["sw_vers", "-productVersion"]),
        "python-zipfile": "%d.%d" % sys.version_info[:2],
    }
    files: list[tuple[str, str, bytes]] = []
    with tempfile.TemporaryDirectory(prefix="docx-producers-", dir=ROOT / "target") as t:
        tmp = Path(t)
        (tmp / "chart.png").write_bytes(png(120, 80, 1, False))
        for name, data in python_docx(tmp).items():
            files.append((name, "python-docx", data))
        (tmp / "doc.md").write_text(MD)
        subprocess.run(["pandoc", "doc.md", "-o", "pandoc.docx"], cwd=tmp, check=True)
        files.append(("pandoc-markdown", "pandoc", (tmp / "pandoc.docx").read_bytes()))
        (tmp / "doc.html").write_text(HTML)
        (tmp / "doc.rtf").write_text(RTF)
        for src in ("doc.html", "doc.rtf"):
            subprocess.run(["textutil", "-convert", "docx", "-output", f"textutil-{src}.docx", src], cwd=tmp, check=True)
            files.append((f"textutil-{src.split('.')[1]}", "textutil", (tmp / f"textutil-{src}.docx").read_bytes()))
        lo = soffice_convert(tmp / "doc.html", tmp, "html")
        if lo:
            files.append(("libreoffice-from-html", "libreoffice", lo))
            for name in ("python-docx-comments-tracked-changes", "pandoc-markdown", "python-docx-large", "python-docx-nested-tables"):
                src = tmp / f"rt-{name}.docx"
                src.write_bytes(next(d for n, _, d in files if n == name))
                files.append((f"libreoffice-resave-{name}", "libreoffice", soffice_convert(src, tmp, f"rt-{name}")))
        java = shutil.which("java")
        if java:
            versions["java-zipoutputstream"] = version([java, "-version"])
            for name in ("python-docx-comments-tracked-changes", "libreoffice-resave-pandoc-markdown", "python-docx-nested-tables"):
                src, dst = tmp / f"j-{name}.docx", tmp / f"java-{name}.docx"
                src.write_bytes(next(d for n, _, d in files if n == name))
                subprocess.run([java, str(Path(__file__).resolve().parent / "JavaRepack.java"), str(src), str(dst)],
                               check=True, capture_output=True, timeout=300)
                files.append((f"java-zipoutputstream-{name}", "java-zipoutputstream", dst.read_bytes()))
        basic = next(d for n, _, d in files if n == "python-docx-basic")
        for tool, label in (("infozip", "zip-infozip"), ("ditto", "zip-ditto"), ("zipfile", "zip-python-zipfile"),
                            ("infozip-comment", "zip-infozip-with-comment")):
            files.append((label, {"zipfile": "python-zipfile", "infozip-comment": "infozip"}.get(tool, tool),
                          repack(tmp, label, basic, tool)))

    records, problems = [], []
    for name, producer, data in files:
        py, rs = gen.run_python(fn, data), gen.run_rust(a.bin, data)
        rec = {"name": name, "kind": "producer", "producer": producer, "bytes": len(data),
               "docx": base64.b64encode(data).decode(), "python": py, "rust": rs}
        if name in EXPECTED_REFUSALS:
            rec["stricter"] = EXPECTED_REFUSALS[name]
            if "error" in py or "error" not in rs:
                problems.append(f"{name}: expected Python ok and a binary refusal, got {py} / {rs}")
        elif "error" in py or py != rs:
            problems.append(f"{name} ({producer}): python {py.get('error', 'ok')} rust {rs.get('error', 'ok')}")
        records.append(rec)
        print(f"{name:45} {producer:15} {len(data):>9} B  python={py.get('error', 'ok')} rust={rs.get('error', 'ok')}")
    if problems:
        sys.exit("producer files not extracted identically:\n  " + "\n  ".join(problems))
    doc = {"version": 1,
           "about": "noevia#981 producer compatibility: DOCX files from real producers (synthetic content), "
                    "Python ocr/docx_text.py vs the docx-text binary. Generated by noevia-rs tools/gen-docx-producers.py.",
           "producers": versions, "cases": records}
    Path(a.out).write_text(json.dumps(doc, ensure_ascii=True, indent=0) + "\n")
    print(f"wrote {a.out}: {len(records)} files")


if __name__ == "__main__":
    main()
